import json
import hashlib
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SOURCE = Path(r"D:\work\cloud-repository\resume\个人简历.docx")


def run_info(run):
    rpr = run._element.rPr
    fonts = {}
    if rpr is not None and rpr.rFonts is not None:
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            value = rpr.rFonts.get(qn(f"w:{key}"))
            if value:
                fonts[key] = value
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": bool(run.underline) if run.underline is not None else None,
        "font_name": run.font.name,
        "font_size_pt": run.font.size.pt if run.font.size else None,
        "font_color": str(run.font.color.rgb) if run.font.color and run.font.color.rgb else None,
        "fonts": fonts,
    }


def para_info(p, index=None):
    pf = p.paragraph_format
    return {
        "index": index,
        "text": p.text,
        "style": p.style.name if p.style else None,
        "alignment": str(p.alignment),
        "left_indent_pt": pf.left_indent.pt if pf.left_indent else None,
        "right_indent_pt": pf.right_indent.pt if pf.right_indent else None,
        "first_line_indent_pt": pf.first_line_indent.pt if pf.first_line_indent else None,
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": str(pf.line_spacing) if pf.line_spacing else None,
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
        "runs": [run_info(r) for r in p.runs],
    }


def main():
    doc = Document(SOURCE)
    data = {
        "source": str(SOURCE),
        "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "paragraphs": [para_info(p, i) for i, p in enumerate(doc.paragraphs)],
        "tables": [],
        "sections": [],
        "headers": [],
        "footers": [],
        "inline_shapes": len(doc.inline_shapes),
    }
    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            cells = []
            for ci, cell in enumerate(row.cells):
                cells.append({
                    "cell": [ri, ci],
                    "text": cell.text,
                    "paragraphs": [para_info(p) for p in cell.paragraphs],
                })
            rows.append(cells)
        data["tables"].append({"index": ti, "style": table.style.name if table.style else None, "rows": rows})
    for si, section in enumerate(doc.sections):
        data["sections"].append({
            "index": si,
            "page_width_in": section.page_width.inches,
            "page_height_in": section.page_height.inches,
            "top_margin_in": section.top_margin.inches,
            "bottom_margin_in": section.bottom_margin.inches,
            "left_margin_in": section.left_margin.inches,
            "right_margin_in": section.right_margin.inches,
            "header_distance_in": section.header_distance.inches,
            "footer_distance_in": section.footer_distance.inches,
        })
        data["headers"].append([para_info(p) for p in section.header.paragraphs])
        data["footers"].append([para_info(p) for p in section.footer.paragraphs])
    with zipfile.ZipFile(SOURCE) as zf:
        data["package_parts"] = [
            {"path": i.filename, "size": i.file_size, "sha256": hashlib.sha256(zf.read(i.filename)).hexdigest()}
            for i in zf.infolist()
        ]
    out = SOURCE.parent / ".resume_work" / "source_inspection.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
