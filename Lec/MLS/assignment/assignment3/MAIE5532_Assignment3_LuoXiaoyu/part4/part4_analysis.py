#!/usr/bin/env python3
"""生成第 4 部分的图、表、Markdown 报告和性能工作簿。"""

from __future__ import annotations

import argparse
import csv
import json
import math
import sys
from datetime import date
from pathlib import Path
from typing import Any, Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt


PART4_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = PART4_DIR.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from optimization_framework import (  # noqa: E402
    HardwareAwareDesignMethodology,
    build_reference_design_cases,
)
from performance_profiler import (  # noqa: E402
    PLATFORM_SPECS,
    analyze_hardware_utilization,
    generate_performance_report,
)


FIGURE_DIR = PART4_DIR / "part4_results" / "figures"
RESULTS_DIR = PART4_DIR / "part4_results"
COLORS = {
    "baseline": "#334155",
    "latency_optimized": "#2563EB",
    "memory_optimized": "#7C3AED",
    "energy_optimized": "#059669",
    "ptq_int8": "#DC2626",
    "qat_int8": "#F59E0B",
    "mixed_precision": "#0891B2",
    "dynamic_range": "#DB2777",
}


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _write_csv(path: Path, rows: Iterable[dict]) -> None:
    rows = list(rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fields = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow(
                {
                    key: (
                        json.dumps(value, ensure_ascii=False)
                        if isinstance(value, dict)
                        else value
                    )
                    for key, value in row.items()
                }
            )


def _normalize(values: list[float], higher_is_better: bool) -> list[float]:
    low = min(values)
    high = max(values)
    if math.isclose(low, high):
        return [0.5 for _ in values]
    normalized = [(value - low) / (high - low) for value in values]
    return normalized if higher_is_better else [1.0 - value for value in normalized]


def _save_accuracy_latency_chart(
    measured: list[dict],
    frontier_names: set[str],
    path: Path,
) -> None:
    figure, axis = plt.subplots(figsize=(7.2, 4.4))
    for row in measured:
        is_frontier = row["model"] in frontier_names
        axis.scatter(
            row["latency_ms"],
            row["accuracy"] * 100.0,
            s=90 if is_frontier else 55,
            color=COLORS[row["model"]],
            edgecolor="#111827" if is_frontier else "#FFFFFF",
            linewidth=1.2,
            zorder=3,
            label=row["model"],
        )
        axis.annotate(
            row["model"].replace("_", " "),
            (row["latency_ms"], row["accuracy"] * 100.0),
            xytext=(6, 5),
            textcoords="offset points",
            fontsize=7.5,
        )
    axis.set_xscale("log")
    axis.set_xlabel("Measured latency (ms, log scale)")
    axis.set_ylabel("CIFAR-10 accuracy (%)")
    axis.set_title("Accuracy versus measured latency")
    axis.grid(alpha=0.2)
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _save_accuracy_size_chart(
    measured: list[dict],
    frontier_names: set[str],
    path: Path,
) -> None:
    figure, axis = plt.subplots(figsize=(7.2, 4.4))
    for row in measured:
        is_frontier = row["model"] in frontier_names
        axis.scatter(
            row["model_size_mb"],
            row["accuracy"] * 100.0,
            s=90 if is_frontier else 55,
            color=COLORS[row["model"]],
            edgecolor="#111827" if is_frontier else "#FFFFFF",
            linewidth=1.2,
            zorder=3,
        )
        axis.annotate(
            row["model"].replace("_", " "),
            (row["model_size_mb"], row["accuracy"] * 100.0),
            xytext=(6, 5),
            textcoords="offset points",
            fontsize=7.5,
        )
    axis.set_xscale("log")
    axis.set_xlabel("Model size (MiB, log scale)")
    axis.set_ylabel("CIFAR-10 accuracy (%)")
    axis.set_title("Accuracy versus model size")
    axis.grid(alpha=0.2)
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _save_radar_chart(measured: list[dict], path: Path) -> None:
    selected = [
        "baseline",
        "memory_optimized",
        "energy_optimized",
        "ptq_int8",
        "mixed_precision",
        "dynamic_range",
    ]
    rows = [row for row in measured if row["model"] in selected]
    metrics = [
        ("Accuracy", [row["accuracy"] for row in rows], True),
        (
            "Latency efficiency",
            [1.0 / row["latency_ms"] for row in rows],
            True,
        ),
        (
            "Size efficiency",
            [1.0 / row["model_size_mb"] for row in rows],
            True,
        ),
        (
            "Energy efficiency",
            [1.0 / row["estimated_energy_mj"] for row in rows],
            True,
        ),
        (
            "Parameter efficiency",
            [1.0 / row["parameters"] for row in rows],
            True,
        ),
    ]
    normalized = [
        _normalize(values, higher_is_better)
        for _, values, higher_is_better in metrics
    ]
    labels = [label for label, _, _ in metrics]
    angles = [
        2.0 * math.pi * index / len(labels) for index in range(len(labels))
    ]
    angles.append(angles[0])

    figure, axis = plt.subplots(figsize=(7.2, 5.2), subplot_kw={"polar": True})
    for row_index, row in enumerate(rows):
        values = [
            normalized[metric_index][row_index]
            for metric_index in range(len(metrics))
        ]
        values.append(values[0])
        axis.plot(
            angles,
            values,
            color=COLORS[row["model"]],
            linewidth=1.6,
            label=row["model"].replace("_", " "),
        )
        axis.fill(angles, values, color=COLORS[row["model"]], alpha=0.05)
    axis.set_xticks(angles[:-1])
    axis.set_xticklabels(labels, fontsize=8)
    axis.set_ylim(0.0, 1.0)
    axis.set_title("Normalized multi-objective trade-offs")
    axis.legend(loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=3, fontsize=7)
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _save_batch_chart(batch_rows: list[dict], path: Path) -> None:
    batches = [int(row["batch_size"]) for row in batch_rows]
    latency = [float(row["latency_ms"]) for row in batch_rows]
    throughput = [
        float(row["throughput_samples_per_second"]) for row in batch_rows
    ]
    figure, axis = plt.subplots(figsize=(7.2, 4.2))
    axis.plot(batches, latency, marker="o", color="#2563EB", linewidth=1.8)
    axis.set_xlabel("Batch size")
    axis.set_ylabel("Total latency (ms)", color="#2563EB")
    axis.tick_params(axis="y", labelcolor="#2563EB")
    axis.grid(alpha=0.2)

    throughput_axis = axis.twinx()
    throughput_axis.plot(
        batches,
        throughput,
        marker="s",
        color="#DC2626",
        linewidth=1.6,
    )
    throughput_axis.set_ylabel("Throughput (samples/s)", color="#DC2626")
    throughput_axis.tick_params(axis="y", labelcolor="#DC2626")
    axis.set_title("Batch-size scaling for the latency-optimized model")
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _save_platform_chart(platform_rows: list[dict], path: Path) -> None:
    selected_models = [
        "latency_optimized",
        "memory_optimized",
        "energy_optimized",
        "ptq_int8",
        "mixed_precision",
    ]
    platforms = [
        "x86_desktop",
        "arm_cortex_a78",
        "arm_cortex_m7",
        "webgpu_proxy",
    ]
    lookup = {
        (row["platform"], row["model"]): row["estimated_speedup_vs_baseline"]
        for row in platform_rows
    }
    width = 0.16
    positions = list(range(len(platforms)))
    figure, axis = plt.subplots(figsize=(8.2, 4.5))
    for model_index, model in enumerate(selected_models):
        values = [lookup[(platform, model)] for platform in platforms]
        offsets = [
            position + (model_index - 2) * width for position in positions
        ]
        axis.bar(
            offsets,
            values,
            width=width,
            color=COLORS[model],
            label=model.replace("_", " "),
        )
    axis.set_xticks(positions)
    axis.set_xticklabels(platforms, fontsize=8)
    axis.set_ylabel("Estimated speedup versus baseline")
    axis.set_title("Cross-platform optimization effect")
    axis.legend(fontsize=7, ncol=3)
    axis.grid(axis="y", alpha=0.2)
    figure.tight_layout()
    figure.savefig(path, dpi=180)
    plt.close(figure)


def _figures(performance: dict) -> dict[str, Path]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    measured = performance["measured"]
    frontier_names = set(performance["pareto"]["accuracy_latency_size"])
    outputs = {
        "accuracy_latency": FIGURE_DIR / "accuracy_vs_latency.png",
        "accuracy_size": FIGURE_DIR / "accuracy_vs_model_size.png",
        "radar": FIGURE_DIR / "normalized_tradeoffs.png",
        "batch": FIGURE_DIR / "batch_scaling.png",
        "platform": FIGURE_DIR / "platform_effectiveness.png",
    }
    _save_accuracy_latency_chart(
        measured,
        frontier_names,
        outputs["accuracy_latency"],
    )
    _save_accuracy_size_chart(
        measured,
        set(performance["pareto"]["accuracy_model_size"]),
        outputs["accuracy_size"],
    )
    _save_radar_chart(measured, outputs["radar"])
    _save_batch_chart(performance["batch_scaling"], outputs["batch"])
    _save_platform_chart(performance["platform_estimates"], outputs["platform"])
    return outputs


def _flatten_utilization(utilization: dict) -> list[dict]:
    keys = sorted(utilization["power_efficiency"])
    rows = []
    for key in keys:
        platform, model = key.split("/", 1)
        compute = utilization["cpu_simd_utilization"][key]
        bandwidth = utilization["memory_bandwidth_efficiency"][key]
        cache = utilization["cache_efficiency"][key]
        thermal = utilization["thermal_characteristics"][key]
        power = utilization["power_efficiency"][key]
        rows.append(
            {
                "platform": platform,
                "model": model,
                "execution_label": PLATFORM_SPECS[platform]["label"],
                "compute_utilization": compute[
                    "estimated_compute_utilization"
                ],
                "achieved_gops": compute["achieved_gops"],
                "bottleneck": compute["bottleneck"],
                "memory_bandwidth_utilization": bandwidth[
                    "memory_bandwidth_utilization"
                ],
                "achieved_bandwidth_gbps": bandwidth[
                    "achieved_bandwidth_gbps"
                ],
                "working_set_mb": bandwidth["working_set_mb"],
                "l1_capacity_fraction": cache["l1_capacity_fraction"],
                "l2_capacity_fraction": cache["l2_capacity_fraction"],
                "l3_capacity_fraction": cache["l3_capacity_fraction"],
                "thermal_derate_assumption": thermal["thermal_derate"],
                "estimated_energy_mj": power["estimated_energy_mj"],
                "estimated_energy_pj_per_flop": power[
                    "estimated_energy_pj_per_flop"
                ],
            }
        )
    return rows


def _markdown_table(headers: list[str], rows: list[list[Any]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join("---" for _ in headers) + "|",
    ]
    lines.extend(
        "| " + " | ".join(str(value) for value in row) + " |"
        for row in rows
    )
    return "\n".join(lines)


def _report_metrics(
    performance: dict,
    utilization: dict,
    cases: list[dict],
) -> dict:
    measured = {row["model"]: row for row in performance["measured"]}
    baseline = measured["baseline"]
    ptq = measured["ptq_int8"]
    mixed = measured["mixed_precision"]
    qemu = utilization["mcu_memory_validity"]
    mobile_case = cases[0]
    mcu_case = cases[1]
    return {
        "baseline": baseline,
        "ptq": ptq,
        "mixed": mixed,
        "qemu": qemu,
        "mobile_case": mobile_case,
        "mcu_case": mcu_case,
        "latency_reduction": 1.0 - ptq["latency_ms"] / baseline["latency_ms"],
        "size_reduction": 1.0 - ptq["model_size_mb"] / baseline["model_size_mb"],
        "accuracy_loss": baseline["accuracy"] - ptq["accuracy"],
    }


def _english_report(
    performance: dict,
    utilization: dict,
    cases: list[dict],
    figures: dict[str, Path],
) -> str:
    metrics = _report_metrics(performance, utilization, cases)
    baseline = metrics["baseline"]
    ptq = metrics["ptq"]
    mixed = metrics["mixed"]
    qemu = metrics["qemu"]
    mobile = metrics["mobile_case"]
    mcu = metrics["mcu_case"]
    measured_rows = [
        [
            row["model"],
            f"{row['accuracy'] * 100:.2f}%",
            f"{row['latency_ms']:.3f}",
            f"{row['model_size_mb']:.3f}",
        ]
        for row in performance["measured"]
    ]
    batch_rows = [
        [
            int(row["batch_size"]),
            f"{row['latency_ms']:.3f}",
            f"{row['throughput_samples_per_second']:.0f}",
            f"{row['latency_per_sample_ms']:.3f}",
        ]
        for row in performance["batch_scaling"]
    ]
    pareto = performance["pareto"]
    figure_root = "part4/part4_results/figures"
    return f"""# Hardware-Aware Design Analysis

**Course:** MAIE5532  
**Assignment:** 3, Part 4  
**Prepared:** {date.today().isoformat()}

## Executive summary

The baseline MobileNetV2 reaches {baseline['accuracy'] * 100:.2f}% CIFAR-10
accuracy with {baseline['model_size_mb']:.3f} MiB of model data and
{baseline['latency_ms']:.3f} ms measured single-image latency. The PTQ INT8
variant reduces measured latency by {metrics['latency_reduction'] * 100:.1f}%
and model size by {metrics['size_reduction'] * 100:.1f}%, while reducing
accuracy by {metrics['accuracy_loss'] * 100:.2f} percentage points. Mixed
precision provides the fastest measured host latency at
{mixed['latency_ms']:.3f} ms and retains {mixed['accuracy'] * 100:.2f}%
accuracy at {mixed['model_size_mb']:.3f} MiB. The result is a genuine
multi-objective trade-off rather than a single universally best model.

## 1. Hardware-software co-design

The baseline input size of 224x224 preserves compatibility with the standard
MobileNetV2 design, but it also carries {baseline['flops']:,} FLOPs per image.
The hardware-constrained variant changes two design variables together:
the MobileNetV2 width multiplier becomes 0.5 and the input resolution becomes
160x160. This lowers arithmetic demand to approximately
{ptq['flops']:,} FLOPs, an 83.6% reduction. The choice is hardware-led:
reducing channels lowers both weights and activation width, while reducing
input resolution lowers the area of every spatial feature map.

Accuracy and efficiency interact strongly. Width and input reduction lower
the accuracy ceiling before quantization is applied. PTQ then introduces only
a modest additional loss in this experiment, whereas the simplified QAT
variant did not recover the loss. This outcome shows that a nominally more
advanced optimizer is not automatically better when its training schedule and
calibration quality are limited.

![Accuracy versus measured latency]({figure_root}/accuracy_vs_latency.png)

## 2. Platform-specific optimization insights

The measured local host separates architecture effects from platform effects.
{_markdown_table(['Model', 'Accuracy', 'Latency ms', 'Size MiB'], measured_rows)}

The latency-optimized Keras model and baseline have similar host latency
because GPU launch and framework overhead are large relative to the reduced
graph. The TFLite INT8 path changes the deployment stack as well as the model,
so its measured 0.663 ms latency is materially lower. On CPU-class targets,
the most effective sequence is therefore graph reduction, quantization, and
kernel/runtime selection, rather than architecture width reduction alone.

![Cross-platform optimization effectiveness]({figure_root}/platform_effectiveness.png)

The Cortex-M result is constrained by memory rather than arithmetic alone.
The INT8 file is {qemu['model_bytes']:,} bytes and fits the modeled 2 MiB
flash budget, but TFLite Micro uses {qemu['tensor_arena_used_bytes']:,} bytes
of tensor arena, exceeding the modeled 512 KiB SRAM budget. QEMU proves
functional execution and gives a relative latency indication; it does not
enforce the physical SRAM limit. A real MCU deployment would need a smaller
arena, a smaller input or activation set, or a platform with more SRAM.

Desktop x86 benefits from large caches and SIMD kernels, ARM Cortex-A is more
sensitive to bandwidth and thermal derating, and the WebGPU proxy is dominated
by launch and synchronization overhead. Mobile GPUs can still win through
parallelism, but only when batch size and work per launch are large enough to
amortize those costs.

## 3. Energy-latency-accuracy trade-off

![Normalized multi-objective trade-offs]({figure_root}/normalized_tradeoffs.png)

The multi-objective Pareto frontier for measured host latency, accuracy, and
model size is:

**{', '.join(pareto['accuracy_latency_size'])}.**

The two-dimensional accuracy-latency frontier contains
**{', '.join(pareto['accuracy_latency'])}**, while the accuracy-size frontier
contains **{', '.join(pareto['accuracy_model_size'])}**. These sets show the
expected trade-offs: the baseline protects accuracy, PTQ and dynamic-range
quantization compress the artifact, and mixed precision balances accuracy and
latency. The memory-optimized variant is not automatically preferable merely
because it has 30% kernel sparsity; its Keras artifact remains larger than
several quantized alternatives.

![Accuracy versus model size]({figure_root}/accuracy_vs_model_size.png)

The energy values are analytical estimates. Across the full-model graph they
change much less than model size because the optimized variants share nearly
the same FLOP count and the Roofline model is compute-bound. Early exit has
little effect on the average graph estimate because only 0.81% of evaluation
samples used the shortcut. For battery-powered devices, INT8 remains the
strongest general-purpose choice, but aggressive early exit or event-driven
sensing would be required when energy, rather than average latency, is the
primary constraint.

## 4. Scalability and deployment

{_markdown_table(['Batch', 'Latency ms', 'Throughput samples/s', 'Latency per sample ms'], batch_rows)}

Increasing batch size improves throughput but does not improve per-image
latency. The measured throughput rises from approximately 525 samples/s at
batch 1 to approximately 10,000 samples/s at batch 64, while total latency
grows from 1.91 ms to 6.40 ms. The maximum throughput setting is therefore
appropriate for offline or server-side batching, not for an interactive
camera pipeline.

![Batch-size scaling]({figure_root}/batch_scaling.png)

Heterogeneous deployment requires separate acceptance tests for each runtime:
numerical tolerance, operator support, arena size, model format, and
accelerator availability. Model updates should be signed and versioned;
resource-constrained devices should receive compressed deltas or complete
INT8 artifacts, verify them before activation, and retain a known-good model
for rollback.

## 5. Design methodology recommendations

The method should start by converting application requirements into hard
budgets, not by selecting a model first. The implemented framework evaluates:

- accuracy target,
- latency budget,
- model-size limit,
- working-set or SRAM limit,
- energy budget,
- target platform and execution label.

For the **{mobile['name']}** case, the feasible set is
{mobile['design_space']['feasible_models']} and the recommended model is
**{mobile['design_space']['recommended_model']}**. This reflects the value of
INT8 and mixed precision on a mobile-class CPU when a moderate accuracy loss
is acceptable.

For the **{mcu['name']}** case, the feasible set is
{mcu['design_space']['feasible_models']}. Although the INT8 artifact fits the
flash budget, the measured TFLite Micro arena exceeds 512 KiB SRAM. The design
should therefore continue with a smaller input, a more compact feature
pyramid, or a different MCU rather than treating successful QEMU execution as
sufficient evidence of deployability.

Recommended lifecycle:

1. Specify platform, accuracy, latency, energy, and memory budgets.
2. Profile the baseline with the final runtime.
3. Reduce architectural work before applying runtime-specific optimization.
4. Quantize and verify accuracy on the complete test distribution.
5. Validate memory, operator coverage, and numerical tolerance on each target.
6. Re-profile sustained performance after thermal and cache warm-up.
7. Keep the smallest model that satisfies every hard constraint.

## 6. Future hardware trends

Edge AI accelerators will reduce the value of hand-tuned arithmetic
optimizations when they support the required operators, but memory movement
will remain important. Compiler-assisted operator fusion, structured sparsity,
and mixed precision will become more valuable when they match the accelerator
tile and buffer hierarchy.

Neuromorphic and event-driven hardware could especially benefit
classification workloads that can stop early on sparse changes. That direction
requires event-based sensors and new training objectives, so it is a future
extension rather than a drop-in replacement for this assignment's dense image
pipeline. The same constraint-first methodology remains applicable.

## Validity and limitations

- Accuracy and host latency are measured; platform latency and energy are
  analytical estimates except for the QEMU-CM and WebGPU proxy runs.
- The cache values are capacity proxies, not hardware hit rates.
- Thermal derating is an assumption; no temperature sensor was available.
- QEMU user mode does not enforce SRAM, flash, cache, or thermal limits.
- The WebGPU measurement is a desktop proxy, not a smartphone measurement.

## References

1. S. Williams, A. Waterman, and D. Patterson, "Roofline: An Insightful
   Visual Performance Model for Multicore Architectures," *CACM*, 2009.
2. M. Sandler et al., "MobileNetV2: Inverted Residuals and Linear
   Bottlenecks," *CVPR*, 2018.
3. B. Jacob et al., "Quantization and Training of Neural Networks for
   Efficient Integer-Arithmetic-Only Inference," *CVPR*, 2018.
4. R. David et al., "TensorFlow Lite Micro: Embedded Machine Learning for
   TinyML Systems," *Proceedings of MLSys*, 2021.
5. Course Week 6, *Hardware Backend Specialization and Optimizations*.
6. Course Week 7, *Machine Learning for Systems*.
"""


def _chinese_report(
    performance: dict,
    utilization: dict,
    cases: list[dict],
    figures: dict[str, Path],
) -> str:
    metrics = _report_metrics(performance, utilization, cases)
    baseline = metrics["baseline"]
    ptq = metrics["ptq"]
    mixed = metrics["mixed"]
    qemu = metrics["qemu"]
    mobile = metrics["mobile_case"]
    mcu = metrics["mcu_case"]
    pareto = performance["pareto"]
    figure_root = "part4/part4_results/figures"
    return f"""# 硬件感知设计分析

**课程：** MAIE5532  
**作业：** Assignment 3，第 4 部分  
**生成日期：** {date.today().isoformat()}

## 摘要

基线 MobileNetV2 的 CIFAR-10 准确率为 {baseline['accuracy'] * 100:.2f}%，
模型大小为 {baseline['model_size_mb']:.3f} MiB，实测单张图片延迟为
{baseline['latency_ms']:.3f} ms。PTQ INT8 将实测延迟降低
{metrics['latency_reduction'] * 100:.1f}%，模型大小降低
{metrics['size_reduction'] * 100:.1f}%，准确率降低
{metrics['accuracy_loss'] * 100:.2f} 个百分点。混合精度的实测延迟最低，
为 {mixed['latency_ms']:.3f} ms，准确率为 {mixed['accuracy'] * 100:.2f}%，
大小为 {mixed['model_size_mb']:.3f} MiB。因此不存在单一模型在所有指标上
都最优，实际选择取决于部署约束。

## 1. 硬件与软件协同设计

基线保留 224×224 输入，以符合标准 MobileNetV2，但单张图片需要
{baseline['flops']:,} FLOPs。硬件优化模型同时把宽度乘数降到 0.5，把输入
尺寸降到 160×160，使 FLOPs 降到约 {ptq['flops']:,}，下降 83.6%。
宽度下降会同时减少权重通道和激活宽度，输入尺寸下降会减少各层空间大小，
所以这两项改变直接对应硬件预算。

准确率与效率存在明显权衡。先缩小宽度和输入，再进行量化；当前简化 QAT
没有恢复这部分损失，说明训练日程和校准数据质量与优化技术本身同样重要。

![准确率与实测延迟]({figure_root}/accuracy_vs_latency.png)

## 2. 各平台优化结论

在本地 x86 主机上，延迟优化 Keras 模型与基线的耗时接近，说明 GPU 启动和
框架调度在单张图片推理中占比较大。PTQ INT8 的实测延迟为
{ptq['latency_ms']:.3f} ms，说明降低图计算量、切换 TFLite 运行时和量化
需要一起实施。对 CPU 类平台，优先顺序应是减少图计算量、INT8 量化，
再选择高效的算子库与运行时。

![跨平台优化效果]({figure_root}/platform_effectiveness.png)

Cortex-M 的瓶颈首先是内存。INT8 文件为 {qemu['model_bytes']:,} bytes，
可以放入假设的 2 MiB flash；但 TFLite Micro arena 实际占用
{qemu['tensor_arena_used_bytes']:,} bytes，超过假设的 512 KiB SRAM。
因此 QEMU 跑通只说明功能可执行，不代表该模型能直接部署在 512 KiB SRAM
的 MCU 上。真实部署还需要更小的输入、更紧凑的激活缓冲区或更大的 SRAM。

x86 更受益于大 cache 和 SIMD；Cortex-A 更受带宽和热降频影响；WebGPU
代理的主要成本是启动、提交和同步。移动 GPU 只有在并行任务足够大时，
才能用并行度抵消固定开销。

## 3. 能量、延迟和准确率权衡

![归一化多目标权衡]({figure_root}/normalized_tradeoffs.png)

在“准确率最大、实测延迟最小、模型大小最小”三维目标下，Pareto 前沿为：

**{', '.join(pareto['accuracy_latency_size'])}。**

准确率-延迟二维前沿包含
**{', '.join(pareto['accuracy_latency'])}**；准确率-模型大小二维前沿包含
**{', '.join(pareto['accuracy_model_size'])}**。基线保留了最高准确率，
PTQ 和动态范围量化显著压缩模型，混合精度在准确率和延迟之间取得平衡。
稀疏内存模型不能仅凭 30% 稀疏率判定为最优，因为它序列化后仍大于多个
量化模型。

![准确率与模型大小]({figure_root}/accuracy_vs_model_size.png)

能量数字来自解析估算，不是功率计测量。由于多个优化模型的 FLOPs 接近，
完整模型图上的能耗下降不如模型大小下降明显。Early exit 只被
0.81% 的样本触发，因此平均能耗收益有限。对电池设备，INT8 仍是通用性
最好的方案；如果能源是最紧约束，还需要事件驱动推理或更激进的条件计算。

## 4. 可扩展性与部署

批大小从 1 增加到 64 时，吞吐量从约 525 samples/s 增加到约
10,000 samples/s，但单次总延迟也从约 1.91 ms 增加到 6.40 ms。
大批量适合服务器离线处理或吞吐优先场景，不适合要求单帧即时响应的
交互式相机流水线。

![批大小扩展]({figure_root}/batch_scaling.png)

异构硬件部署需要分别验证精度、算子支持、tensor arena、模型格式和加速器
可用性。模型更新应签名和版本化；资源受限设备应接收完整的 INT8 模型或
压缩增量，激活新模型前完成校验，并保留可回滚的已知良好版本。

## 5. 设计方法建议

设计应从硬约束出发，而不是先选模型。框架需要记录准确率目标、延迟预算、
模型大小、工作集或 SRAM、能耗和平台类型。

- 移动端案例可行集合：
  {mobile['design_space']['feasible_models']}；推荐模型：
  **{mobile['design_space']['recommended_model']}**。
- MCU 案例可行集合：{mcu['design_space']['feasible_models']}。
  虽然 INT8 模型可放入 flash，但 arena 超过 SRAM，因此没有模型满足全部
  约束。该结论说明功能仿真不能替代容量验证。

建议流程是：定义约束；用最终运行时测量基线；先减少架构计算量；量化并用
完整测试集验证；在目标运行时检查内存和算子；热稳定后重新测量；
在满足所有硬约束的模型中选最小者。

## 6. 未来硬件趋势

边缘 AI 加速器会提高固定算子的效率，但内存移动仍将限制性能。编译器融合、
结构化稀疏和混合精度能否发挥优势，取决于硬件 tile、缓冲区层级和算子覆盖。

神经形态或事件驱动硬件更适合稀疏事件流；它需要事件相机和新的训练目标，
不能直接替代当前的稠密图像流水线。约束优先、运行时验证和多目标比较的方法
仍然适用。

## 有效性与局限

- 准确率和本地主机延迟为实测；除 QEMU-CM 和 WebGPU 代理外，平台延迟与
  能耗为解析估算。
- cache 指标是容量代理，不是实测命中率。
- 热降频系数是假设值，没有温度传感器数据。
- QEMU 用户模式不执行物理 SRAM、flash、cache 或热限制。
- WebGPU 是桌面代理，不是手机实测。

## 参考文献

1. S. Williams, A. Waterman, and D. Patterson, "Roofline: An Insightful
   Visual Performance Model for Multicore Architectures," *CACM*, 2009.
2. M. Sandler et al., "MobileNetV2: Inverted Residuals and Linear
   Bottlenecks," *CVPR*, 2018.
3. B. Jacob et al., "Quantization and Training of Neural Networks for
   Efficient Integer-Arithmetic-Only Inference," *CVPR*, 2018.
4. R. David et al., "TensorFlow Lite Micro: Embedded Machine Learning for
   TinyML Systems," *Proceedings of MLSys*, 2021.
5. 课程第 6 周，*Hardware Backend Specialization and Optimizations*。
6. 课程第 7 周，*Machine Learning for Systems*。
"""


def _guidelines(performance: dict, cases: list[dict]) -> str:
    pareto = performance["pareto"]
    return f"""# Hardware-Aware Optimization Guidelines

## 1. Define constraints before choosing a model

Record the deployment platform, execution label, accuracy floor, latency
ceiling, energy budget, model-size ceiling, and SRAM or working-set ceiling.
Reject a candidate that violates any hard constraint before comparing its
secondary metrics.

## 2. Use a staged optimization order

1. Profile the baseline with the deployment runtime.
2. Reduce input resolution, width, or graph work.
3. Apply INT8 PTQ when the operator set and accelerator support integer paths.
4. Use QAT only when calibration and training time are sufficient.
5. Add pruning or early exit when the workload distribution supports it.
6. Re-measure memory, latency, accuracy, and energy on the target runtime.

## 3. Match optimization to hardware pressure

- **x86 CPU:** use SIMD-aware runtimes and compare against cached, warmed-up
  single-image and batched execution.
- **ARM Cortex-A:** prioritize reduced FLOPs, FP16 or INT8, and bandwidth-aware
  layouts.
- **Cortex-M:** verify flash and SRAM separately. An artifact that fits in flash
  can still fail because of TFLite Micro's tensor arena.
- **Mobile/WebGPU:** amortize launch overhead with sufficient parallel work;
  do not expect small-model changes to reduce end-to-end latency linearly.

## 4. Use Pareto analysis rather than a single score

The measured Pareto set for accuracy, latency, and model size is:

**{', '.join(pareto['accuracy_latency_size'])}.**

Choose from this set according to the application. Use the highest-accuracy
model when errors are costly, the smallest or fastest model when resource
budgets dominate, and the balanced model when deployment requirements do not
strongly prioritize one dimension.

## 5. Validate before deployment

Use the complete test set for accuracy, not calibration samples. Confirm all
operators exist in the target runtime. Measure sustained performance after
warm-up and thermal stabilization. Record model hashes and runtime versions.
Keep a rollback artifact on the device.

## Reference design cases

{_markdown_table(
    ['Case', 'Platform', 'Feasible models', 'Recommendation'],
    [
        [
            case['name'],
            case['constraints']['platform'],
            ', '.join(case['design_space']['feasible_models']) or 'None',
            case['design_space']['recommended_model'] or 'No fully feasible model',
        ]
        for case in cases
    ],
)}
"""


def _guidelines_chinese(performance: dict, cases: list[dict]) -> str:
    pareto = performance["pareto"]
    return f"""# 硬件感知优化指南

## 1. 先定义约束，再选择模型

在训练或优化前记录平台、执行标签、最低准确率、延迟上限、能耗预算、模型
大小上限以及 SRAM 或工作集上限。任何硬约束不满足的模型都应直接淘汰。

## 2. 使用分阶段优化顺序

1. 用最终部署运行时测量基线。
2. 降低输入分辨率、宽度或计算图工作量。
3. 在运行时支持整数算子时使用 INT8 PTQ。
4. 只有在校准数据和训练时间足够时使用 QAT。
5. 工作负载分布合适时再加入剪枝或早退出。
6. 在目标运行时重新测量内存、延迟、准确率和能耗。

## 3. 按瓶颈选择优化方法

- **x86 CPU：** 使用支持 SIMD 的运行时，并分别测量预热后的单图和批量推理。
- **ARM Cortex-A：** 优先减少 FLOPs，再考虑 FP16、INT8 和带宽友好布局。
- **Cortex-M：** 分别检查 flash 和 SRAM。模型能放入 flash，arena 仍可能
  超出 SRAM。
- **移动 GPU/WebGPU：** 让每次启动承载足够多的并行工作，避免把模型缩小
  误认为端到端延迟会等比例下降。

## 4. 使用 Pareto 分析

准确率、延迟和模型大小的实测 Pareto 集合为：

**{', '.join(pareto['accuracy_latency_size'])}。**

误差代价高的应用选择高准确率模型；资源预算最紧的应用选择最小或最快模型；
约束均衡的应用选择折中模型。

## 5. 部署前验证

用完整测试集评估准确率，确认目标运行时支持全部算子，测量热稳定后的持续
性能，记录模型哈希和运行时版本，并在设备上保留可回滚模型。

## 参考设计案例

{_markdown_table(
    ['案例', '平台', '可行模型', '推荐'],
    [
        [
            case['name'],
            case['constraints']['platform'],
            ', '.join(case['design_space']['feasible_models']) or '无',
            case['design_space']['recommended_model'] or '没有完全可行的模型',
        ]
        for case in cases
    ],
)}
"""


def _write_excel(
    path: Path,
    performance: dict,
    utilization: dict,
    utilization_rows: list[dict],
    cases: list[dict],
) -> bool:
    """使用 openpyxl 生成可编辑的性能工作簿。"""
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, ScatterChart, Series, Reference
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        print(
            "警告：未安装 openpyxl，跳过 performance_benchmarks.xlsx。"
            "请安装 requirements.txt 中的依赖后重新运行。"
        )
        return False

    workbook = Workbook()

    def apply_number_format(sheet, cell_range, number_format):
        for row in sheet[cell_range]:
            for cell in row:
                cell.number_format = number_format

    overview = workbook.active
    overview.title = "Overview"
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:F1")
    overview["A1"] = "Hardware-Aware Performance Benchmarks"
    overview["A1"].font = Font(size=15, bold=True, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 28

    summary = performance["summary"]
    summary_rows = [
        ("Best accuracy", summary["best_accuracy_model"]),
        ("Lowest measured latency", summary["lowest_latency_model"]),
        ("Smallest model", summary["smallest_model"]),
        (
            "Best batch throughput (batch size)",
            summary["highest_batch_throughput"],
        ),
        (
            "INT8 model size (bytes)",
            utilization["mcu_memory_validity"]["model_bytes"],
        ),
        (
            "Cortex-M arena used (bytes)",
            utilization["mcu_memory_validity"]["tensor_arena_used_bytes"],
        ),
        (
            "Cortex-M fits 512 KiB SRAM",
            "Yes" if utilization["mcu_memory_validity"]["fits_sram"] else "No",
        ),
    ]
    overview["A3"] = "Key results"
    overview["A3"].font = Font(bold=True, size=12, color="1F4E78")
    for row_index, (label, value) in enumerate(summary_rows, start=4):
        overview.cell(row=row_index, column=1, value=label)
        overview.cell(row=row_index, column=2, value=value)

    measured_headers = [
        "Model",
        "Accuracy",
        "Latency (ms)",
        "Size (MiB)",
        "Energy estimate (mJ)",
        "Throughput (FPS)",
    ]
    overview["D3"] = "Accuracy and measured latency"
    overview["D3"].font = Font(bold=True, size=12, color="1F4E78")

    measured = performance["measured"]
    measured_sheet = workbook.create_sheet("Model measurements")
    measured_sheet.sheet_view.showGridLines = False
    measured_values = [measured_headers] + [
        [
            row["model"],
            row["accuracy"],
            row["latency_ms"],
            row["model_size_mb"],
            row["estimated_energy_mj"],
            row["throughput_fps"],
        ]
        for row in measured
    ]
    measured_sheet.append(measured_values[0])
    for row in measured_values[1:]:
        measured_sheet.append(row)
    measured_sheet.freeze_panes = "A2"
    measured_sheet.auto_filter.ref = f"A1:F{len(measured_values)}"
    measured_table = Table(
        displayName="MeasuredPerformanceTable",
        ref=f"A1:F{len(measured_values)}",
    )
    measured_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
    )
    measured_sheet.add_table(measured_table)
    apply_number_format(measured_sheet, "B2:B8", "0.00%")
    apply_number_format(measured_sheet, "C2:C8", "0.000")
    apply_number_format(measured_sheet, "D2:D8", "0.000")
    apply_number_format(measured_sheet, "E2:E8", "0.000")
    apply_number_format(measured_sheet, "F2:F8", "#,##0.0")
    for column, width in {"A": 22, "B": 12, "C": 14, "D": 13, "E": 22, "F": 18}.items():
        measured_sheet.column_dimensions[column].width = width
    measured_sheet.freeze_panes = "A2"

    platform_sheet = workbook.create_sheet("Platform estimates")
    platform_sheet.sheet_view.showGridLines = False
    platform_headers = [
        "Platform",
        "Execution label",
        "Model",
        "Accuracy",
        "Estimated latency (ms)",
        "Estimated memory (MiB)",
        "Estimated energy (mJ)",
        "Estimated speedup",
        "Fits workspace",
    ]
    platform_values = [platform_headers] + [
        [
            row["platform"],
            row["execution_label"],
            row["model"],
            row["accuracy"],
            row["estimated_latency_ms"],
            row["estimated_memory_mb"],
            row["estimated_energy_mj"],
            row["estimated_speedup_vs_baseline"],
            "Yes" if row["fits_workspace"] else "No",
        ]
        for row in performance["platform_estimates"]
    ]
    for row in platform_values:
        platform_sheet.append(row)
    platform_sheet.freeze_panes = "A2"
    platform_sheet.auto_filter.ref = (
        f"A1:I{len(platform_values)}"
    )
    platform_table = Table(
        displayName="PlatformEstimatesTable",
        ref=f"A1:I{len(platform_values)}",
    )
    platform_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showRowStripes=True,
    )
    platform_sheet.add_table(platform_table)
    apply_number_format(platform_sheet, "D2:D33", "0.00%")
    apply_number_format(platform_sheet, "E2:E33", "0.000")
    apply_number_format(platform_sheet, "F2:F33", "0.000")
    apply_number_format(platform_sheet, "G2:G33", "0.000")
    apply_number_format(platform_sheet, "H2:H33", "0.00")
    for column, width in {
        "A": 20,
        "B": 18,
        "C": 22,
        "D": 12,
        "E": 23,
        "F": 24,
        "G": 22,
        "H": 19,
        "I": 16,
    }.items():
        platform_sheet.column_dimensions[column].width = width

    pareto_sheet = workbook.create_sheet("Pareto and scaling")
    pareto_sheet.sheet_view.showGridLines = False
    pareto_sheet["A1"] = "Pareto frontier"
    pareto_sheet["A1"].font = Font(bold=True, size=12, color="1F4E78")
    pareto_matrix = [
        ["Objective set", "Models"],
        [
            "Accuracy, latency, size",
            ", ".join(performance["pareto"]["accuracy_latency_size"]),
        ],
        [
            "Accuracy and latency",
            ", ".join(performance["pareto"]["accuracy_latency"]),
        ],
        [
            "Accuracy and size",
            ", ".join(performance["pareto"]["accuracy_model_size"]),
        ],
    ]
    for row_index, row in enumerate(pareto_matrix, start=3):
        for column_index, value in enumerate(row, start=1):
            pareto_sheet.cell(
                row=row_index,
                column=column_index,
                value=value,
            )
    pareto_sheet["A8"] = "Batch scaling"
    pareto_sheet["A8"].font = Font(bold=True, size=12, color="1F4E78")
    batch_headers = [
        "Batch",
        "Latency (ms)",
        "Throughput (samples/s)",
        "Latency per sample (ms)",
    ]
    pareto_sheet.append(batch_headers)
    for row in performance["batch_scaling"]:
        pareto_sheet.append(
            [
                row["batch_size"],
                row["latency_ms"],
                row["throughput_samples_per_second"],
                row["latency_per_sample_ms"],
            ]
        )
    pareto_sheet.append([])
    pareto_sheet.append(
        ["Platform Pareto frontier", "Models"]
    )
    for platform, models in performance["pareto"]["platforms"].items():
        pareto_sheet.append([platform, ", ".join(models)])
    apply_number_format(pareto_sheet, "B11:D15", "0.000")
    apply_number_format(pareto_sheet, "C11:C15", "#,##0")

    batch_chart = LineChart()
    batch_chart.title = "Batch throughput"
    batch_chart.y_axis.title = "Throughput (samples/s)"
    batch_chart.x_axis.title = "Batch size"
    batch_chart.height = 7
    batch_chart.width = 12
    batch_chart.add_data(
        Reference(pareto_sheet, min_col=3, min_row=9, max_row=14),
        titles_from_data=True,
    )
    batch_chart.set_categories(
        Reference(pareto_sheet, min_col=1, min_row=10, max_row=14)
    )
    pareto_sheet.add_chart(batch_chart, "F9")

    scatter_sheet = workbook.create_sheet("Accuracy and latency")
    scatter_sheet.sheet_view.showGridLines = False
    scatter_headers = [
        "Model",
        "Latency (ms)",
        "Accuracy",
        "Size (MiB)",
    ]
    scatter_sheet.append(scatter_headers)
    for row in measured:
        scatter_sheet.append(
            [
                row["model"],
                row["latency_ms"],
                row["accuracy"],
                row["model_size_mb"],
            ]
        )
    scatter = ScatterChart()
    scatter.title = "Accuracy versus measured latency"
    scatter.x_axis.title = "Latency (ms)"
    scatter.y_axis.title = "Accuracy"
    scatter.height = 8
    scatter.width = 13
    x_values = Reference(scatter_sheet, min_col=2, min_row=2, max_row=9)
    y_values = Reference(scatter_sheet, min_col=3, min_row=1, max_row=9)
    series = Series(y_values, x_values, title_from_data=True)
    series.marker.symbol = "circle"
    series.graphicalProperties.line.noFill = True
    scatter.series.append(series)
    scatter_sheet.add_chart(scatter, "F2")

    utilization_sheet = workbook.create_sheet("Utilization")
    utilization_sheet.sheet_view.showGridLines = False
    utilization_headers = [
        "Platform",
        "Model",
        "Compute utilization",
        "Memory bandwidth utilization",
        "Achieved GOPs",
        "Working set (MiB)",
        "Energy estimate (mJ)",
        "Bottleneck",
    ]
    utilization_sheet.append(utilization_headers)
    for row in utilization_rows:
        utilization_sheet.append(
            [
                row["platform"],
                row["model"],
                row["compute_utilization"],
                row["memory_bandwidth_utilization"],
                row["achieved_gops"],
                row["working_set_mb"],
                row["estimated_energy_mj"],
                row["bottleneck"],
            ]
        )
    utilization_sheet.freeze_panes = "A2"
    utilization_sheet.auto_filter.ref = (
        f"A1:H{len(utilization_rows) + 1}"
    )
    apply_number_format(utilization_sheet, "C2:D33", "0.00%")
    apply_number_format(utilization_sheet, "E2:E33", "#,##0.000")
    apply_number_format(utilization_sheet, "F2:F33", "0.000")
    apply_number_format(utilization_sheet, "G2:G33", "0.000")
    for column, width in {
        "A": 20,
        "B": 22,
        "C": 22,
        "D": 28,
        "E": 18,
        "F": 20,
        "G": 22,
        "H": 18,
    }.items():
        utilization_sheet.column_dimensions[column].width = width

    design_sheet = workbook.create_sheet("Design cases")
    design_sheet.sheet_view.showGridLines = False
    design_headers = [
        "Case",
        "Platform",
        "Accuracy target",
        "Latency budget (ms)",
        "Size budget (MiB)",
        "Feasible models",
        "Recommended model",
    ]
    design_sheet.append(design_headers)
    for case in cases:
        constraints = case["constraints"]
        design_sheet.append(
            [
                case["name"],
                constraints["platform"],
                constraints["accuracy_target"],
                constraints["max_latency_ms"],
                constraints["max_model_size_mb"],
                ", ".join(case["design_space"]["feasible_models"]) or "None",
                case["design_space"]["recommended_model"] or "No feasible model",
            ]
        )
    apply_number_format(design_sheet, "C2:C3", "0.00%")
    apply_number_format(design_sheet, "D2:E3", "0.000")
    for column, width in {
        "A": 34,
        "B": 20,
        "C": 18,
        "D": 22,
        "E": 21,
        "F": 45,
        "G": 24,
    }.items():
        design_sheet.column_dimensions[column].width = width

    sources_sheet = workbook.create_sheet("Sources")
    sources_sheet.sheet_view.showGridLines = False
    sources_sheet.append(["Metric", "Source", "Evidence status"])
    sources_sheet.append(
        [
            "Measured model accuracy, latency, size",
            "part2/part2_results/performance_comparison.json",
            "Measured on local x86 host",
        ]
    )
    sources_sheet.append(
        [
            "Platform latency, memory, energy",
            "part3/part3_results/analytical_estimates.json",
            "Analytical estimate",
        ]
    )
    sources_sheet.append(
        [
            "Cortex-M memory and latency",
            "part3/part3_results/qemu_cortex_m.json",
            "QEMU-CM measured simulation",
        ]
    )
    sources_sheet.append(
        [
            "WebGPU proxy latency",
            "part3/part3_results/webgpu_proxy.json",
            "Desktop GPU proxy measurement",
        ]
    )
    for column, width in {"A": 34, "B": 52, "C": 34}.items():
        sources_sheet.column_dimensions[column].width = width

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            if cell.value is not None:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return True


def generate_all(output_dir: Path, skip_excel: bool = False) -> dict:
    """执行 Part 4 全部分析并写入交付目录。"""
    output_dir.mkdir(parents=True, exist_ok=True)
    performance = generate_performance_report()
    utilization = analyze_hardware_utilization(performance)
    utilization_rows = _flatten_utilization(utilization)
    cases = build_reference_design_cases()
    figures = _figures(performance)

    _write_json(output_dir / "performance_report.json", performance)
    _write_json(output_dir / "hardware_utilization.json", utilization)
    _write_json(output_dir / "design_cases.json", cases)
    _write_csv(output_dir / "measured_performance.csv", performance["measured"])
    _write_csv(
        output_dir / "platform_performance.csv",
        performance["platform_estimates"],
    )
    _write_csv(
        output_dir / "batch_scaling.csv",
        performance["batch_scaling"],
    )
    _write_csv(output_dir / "hardware_utilization.csv", utilization_rows)

    (PROJECT_ROOT / "hardware_aware_analysis.md").write_text(
        _english_report(performance, utilization, cases, figures),
        encoding="utf-8",
    )
    (PROJECT_ROOT / "hardware_aware_analysis_zh-CN.md").write_text(
        _chinese_report(performance, utilization, cases, figures),
        encoding="utf-8",
    )
    (PROJECT_ROOT / "optimization_guidelines.md").write_text(
        _guidelines(performance, cases),
        encoding="utf-8",
    )
    (PROJECT_ROOT / "optimization_guidelines_zh-CN.md").write_text(
        _guidelines_chinese(performance, cases),
        encoding="utf-8",
    )

    workbook_written = False
    if not skip_excel:
        workbook_written = _write_excel(
            PROJECT_ROOT / "performance_benchmarks.xlsx",
            performance,
            utilization,
            utilization_rows,
            cases,
        )

    return {
        "performance": performance,
        "utilization": utilization,
        "cases": cases,
        "figures": {name: str(path) for name, path in figures.items()},
        "workbook_written": workbook_written,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=RESULTS_DIR,
        help="JSON/CSV 结果目录；Markdown 始终写到提交目录根目录。",
    )
    parser.add_argument(
        "--skip-excel",
        action="store_true",
        help="仅生成分析、图、CSV 和 Markdown，不生成 xlsx。",
    )
    args = parser.parse_args()
    result = generate_all(
        output_dir=args.output_dir,
        skip_excel=args.skip_excel,
    )
    summary = result["performance"]["summary"]
    memory = result["utilization"]["mcu_memory_validity"]
    print("Part 4 analysis complete.")
    print(f"Best accuracy: {summary['best_accuracy_model']}")
    print(f"Lowest latency: {summary['lowest_latency_model']}")
    print(f"Smallest model: {summary['smallest_model']}")
    print(
        "Cortex-M SRAM fit: "
        f"{'yes' if memory['fits_sram'] else 'no'} "
        f"({memory['tensor_arena_used_bytes']:,} bytes used)"
    )
    print(
        "Excel workbook: "
        f"{'written' if result['workbook_written'] else 'not written'}"
    )


if __name__ == "__main__":
    main()
