from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\work\cloud-repository\resume")
SOURCE = ROOT / ".resume_work" / "source_copy.docx"
ZH_OUT = ROOT / "罗霄羽_中文简历.docx"
EN_OUT = ROOT / "Xiaoyu_Luo_Resume_EN.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
DARK = RGBColor(0x32, 0x32, 0x32)
GRAY = RGBColor(0x54, 0x54, 0x54)


def set_font(run, name, size, color=GRAY, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)


def set_east_asian_layout(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    snap = ppr.find(qn("w:snapToGrid"))
    if snap is None:
        snap = OxmlElement("w:snapToGrid")
        ppr.append(snap)
    snap.set(qn("w:val"), "0")


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def configure_styles(doc, font_name):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = font_name
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = GRAY
    normal._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if "Title" not in styles:
        title = styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    else:
        title = styles["Title"]
    title.font.name = font_name
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)

    if "Resume Section" not in styles:
        section = styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
    else:
        section = styles["Resume Section"]
    section.font.name = font_name
    section.font.size = Pt(12.5)
    section.font.bold = True
    section.font.color.rgb = DARK
    section._element.rPr.rFonts.set(qn("w:ascii"), font_name)
    section._element.rPr.rFonts.set(qn("w:hAnsi"), font_name)
    section._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    section.paragraph_format.space_before = Pt(8)
    section.paragraph_format.space_after = Pt(2.5)
    section.paragraph_format.keep_with_next = True


def add_title(doc, title_text, contact_text, font_name):
    p = doc.add_paragraph(style="Title")
    set_east_asian_layout(p)
    set_font(p.add_run(title_text), font_name, 24, BLACK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    set_east_asian_layout(p)
    set_font(p.add_run(contact_text), font_name, 9.2, GRAY)


def add_section(doc, text, font_name):
    p = doc.add_paragraph(style="Resume Section")
    set_east_asian_layout(p)
    set_font(p.add_run(text), font_name, 12.5, DARK, bold=True)
    return p


def add_entry(doc, title, date, font_name, role=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0.8)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.12), WD_TAB_ALIGNMENT.RIGHT)
    set_east_asian_layout(p)
    set_font(p.add_run(title), font_name, 10, DARK, bold=True)
    if role:
        set_font(p.add_run(" | " + role), font_name, 10, DARK, bold=False)
    set_font(p.add_run("\t" + date), font_name, 9.7, DARK)
    return p


def add_detail(doc, text, font_name, bold_label=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.02)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.06
    set_east_asian_layout(p)
    if bold_label and text.startswith(bold_label):
        set_font(p.add_run(bold_label), font_name, 9.8, GRAY, bold=True)
        set_font(p.add_run(text[len(bold_label):]), font_name, 9.8, GRAY)
    else:
        set_font(p.add_run(text), font_name, 9.8, GRAY)
    return p


def add_bullet(doc, text, font_name):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.06
    set_east_asian_layout(p)
    set_font(p.add_run("• "), font_name, 9.6, GRAY)
    set_font(p.add_run(text), font_name, 9.6, GRAY)
    return p


def finalize(doc, title, author, language):
    sec = doc.sections[0]
    sec.page_width = Inches(8.268)
    sec.page_height = Inches(11.693)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(1.055)
    sec.right_margin = Inches(1.055)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.3)
    doc.core_properties.title = title
    doc.core_properties.author = author
    settings = doc.settings._element
    lang = settings.find(qn("w:themeFontLang"))
    if lang is not None:
        lang.set(qn("w:val"), language)


def build_chinese():
    font = "宋体"
    doc = Document(SOURCE)
    clear_body(doc)
    configure_styles(doc, font)
    add_title(doc, "个人简历", "罗霄羽  |  +86 13631635573  |  +852 91967880  |  13310831343@163.com", font)

    add_section(doc, "教育背景", font)
    add_entry(doc, "香港科技大学", "2026-09 ~ 至今", font, "人工智能与创业（AIE）· 全日制研究生在读")
    add_entry(doc, "华南师范大学", "2022-09 ~ 2026-07", font, "软件工程 · 本科")
    add_detail(doc, "GPA：3.33", font, "GPA：")

    add_section(doc, "项目经验", font)
    add_entry(doc, "机场跑道调度优化算法设计", "2026", font, "本科毕业设计")
    add_bullet(doc, "基于蚁群算法设计跑道调度优化方案，在跑道数量、最短安全间隔和计划起降时间等约束下，以最小化总体航班延迟为目标生成优化调度表。", font)
    add_bullet(doc, "采用前后端架构：前端采集跑道数、安全间隔和计划航班起降时间等参数，并完成有效值与合规性校验。", font)
    add_bullet(doc, "后端运行优化算法并将调度结果返回前端，展示优化后的航班起降顺序与调度建议。", font)

    add_section(doc, "实习经历", font)
    add_entry(doc, "北京雁翎网卫智能科技有限公司", "2025-07 ~ 2025-09", font, "二进制模糊测试 Agent 开发实习生")
    add_bullet(doc, "负责面向二进制可执行程序的自动化模糊测试 Agent 开发，学习并应用污点分析、模糊测试和 AFL 工作流。", font)
    add_bullet(doc, "围绕初始样例生成、种子变异、结果分析和报告生成等环节探索 LLM 集成，推进模糊测试流程自动化。", font)
    add_entry(doc, "广州网律互联网科技有限公司", "2024-10 ~ 2024-12", font, "服务器运维实习生")
    add_bullet(doc, "负责服务器运维和后端优化，编写堡垒机数据迁移脚本，并参与后端服务的维护与调试。", font)
    add_bullet(doc, "使用 FastAPI、Nginx 和 Linux 完成后端服务搭建、代理配置与脚本编写，并参与需求分析、代码评审和部署流程。", font)

    add_section(doc, "技能特长", font)
    add_detail(doc, "编程与数据：Python、Java、C、SQL、MySQL；能够阅读并编写基础汇编代码。", font, "编程与数据：")
    add_detail(doc, "后端与运维：FastAPI、Spring Boot、Linux、Nginx、Shell 脚本。", font, "后端与运维：")
    add_detail(doc, "Web 开发：HTML、CSS、JavaScript、Vue；具备前后端协作与数据校验经验。", font, "Web 开发：")
    add_detail(doc, "算法与安全：蚁群算法、AFL 模糊测试、污点分析。", font, "算法与安全：")
    add_detail(doc, "语言能力：雅思 7.0，可使用英语阅读技术资料并进行工作交流。", font, "语言能力：")

    finalize(doc, "罗霄羽 中文简历", "罗霄羽", "zh-CN")
    doc.save(ZH_OUT)


def build_english():
    font = "Arial"
    doc = Document(SOURCE)
    clear_body(doc)
    configure_styles(doc, font)
    add_title(doc, "RESUME", "Xiaoyu Luo  |  +86 13631635573  |  +852 91967880  |  13310831343@163.com", font)

    add_section(doc, "EDUCATION", font)
    add_entry(doc, "The Hong Kong University of Science and Technology", "Sep 2026 - Present", font)
    add_detail(doc, "Full-time Postgraduate Student, Artificial Intelligence and Entrepreneurship (AIE)", font)
    add_entry(doc, "South China Normal University", "Sep 2022 - Jul 2026", font)
    add_detail(doc, "Bachelor's Degree in Software Engineering | GPA: 3.33", font)

    add_section(doc, "PROJECT EXPERIENCE", font)
    add_entry(doc, "Airport Runway Scheduling Optimization Algorithm", "2026", font, "Undergraduate Capstone Project")
    add_bullet(doc, "Designed an ant colony optimization approach that generates an improved aircraft arrival and departure schedule while minimizing total flight delay under runway-count, minimum-separation, and planned-time constraints.", font)
    add_bullet(doc, "Built a front-end/back-end workflow in which users enter airport and flight parameters; the front end validates value ranges and compliance before submission.", font)
    add_bullet(doc, "Implemented the back-end optimization flow and returned the optimized sequence and scheduling recommendations for display in the web interface.", font)

    add_section(doc, "INTERNSHIP EXPERIENCE", font)
    add_entry(doc, "Beijing Yanling Wangwei Intelligent Technology Co., Ltd.", "Jul 2025 - Sep 2025", font)
    add_detail(doc, "Binary Fuzzing Agent Development Intern", font)
    add_bullet(doc, "Developed an automated fuzzing agent for binary executables and gained hands-on experience with taint analysis, fuzz testing, and AFL workflows.", font)
    add_bullet(doc, "Explored LLM integration across initial test-case generation, seed mutation, result analysis, and report generation to automate the fuzzing workflow.", font)
    add_entry(doc, "Guangzhou Wanglv Internet Technology Co., Ltd.", "Oct 2024 - Dec 2024", font)
    add_detail(doc, "Server Operations Intern", font)
    add_bullet(doc, "Maintained servers and optimized back-end services, including development of a bastion-host data migration script and support for service debugging.", font)
    add_bullet(doc, "Used FastAPI, Nginx, Linux, and shell scripting for service setup and operations, and participated in requirements analysis, code review, and deployment.", font)

    add_section(doc, "TECHNICAL SKILLS", font)
    add_detail(doc, "Programming and Data: Python, Java, C, SQL, MySQL; basic assembly reading and programming.", font, "Programming and Data: ")
    add_detail(doc, "Back End and Operations: FastAPI, Spring Boot, Linux, Nginx, shell scripting.", font, "Back End and Operations: ")
    add_detail(doc, "Web: HTML, CSS, JavaScript, Vue; front-end/back-end integration and input validation.", font, "Web: ")
    add_detail(doc, "Algorithms and Security: ant colony optimization, AFL-based fuzzing, and taint analysis.", font, "Algorithms and Security: ")
    add_detail(doc, "Languages: Mandarin Chinese (native); English (IELTS 7.0).", font, "Languages: ")

    finalize(doc, "Xiaoyu Luo Resume", "Xiaoyu Luo", "en-US")
    doc.save(EN_OUT)


if __name__ == "__main__":
    build_chinese()
    build_english()
    print(ZH_OUT)
    print(EN_OUT)
